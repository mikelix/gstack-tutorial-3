# -*- coding: utf-8 -*-
"""McKinsey-style slide primitives for gstack Tutorial No.3.

Palette, typography and layout follow MCKINSEY_DOCX_PLAYBOOK.docx (the
house style used for the gstack tutorial series), extended from documents
to slides.

The defining convention here is the ACTION TITLE: every content slide's
title is a full sentence stating the takeaway ("Four gates pass mechanically;
the fifth needs a human"), not a topic label ("Gates"). A reader flipping
only the titles must get the whole argument. That single rule is most of
what separates a consulting deck from a corporate one.

Design rules enforced by these helpers:
  - 16:9, 13.333 x 7.5in, consistent 0.75in left/right margin grid
  - Action title top-left, thin navy rule beneath, takeaway-first
  - Generous whitespace; no drop shadows, no gradients, no clip art
  - Source/footnote line bottom-left, page number bottom-right
  - Every text run carries an explicit eastAsia font (CJK rule from the
    playbook Section 6 — without it Word/PowerPoint substitute a serif
    for Chinese glyphs on some builds)
"""
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement

# ---------------------------------------------------------------- palette
# Straight from MCKINSEY_DOCX_PLAYBOOK.docx Color Palette table.
NAVY = RGBColor(0x1F, 0x4E, 0x78)   # primary: headings, rules, key numbers
NAVY_DK = RGBColor(0x14, 0x35, 0x52)  # deeper navy for full-bleed dividers
BLACK = RGBColor(0x00, 0x00, 0x00)   # body text
CAPTION = RGBColor(0x59, 0x59, 0x59)  # subtitles, labels
MUTED = RGBColor(0x80, 0x80, 0x80)   # dates, footnotes, page numbers
FILL = RGBColor(0xF2, 0xF2, 0xF2)   # code blocks, zebra rows
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
RULE = RGBColor(0xD8, 0xDE, 0xE4)   # hairline separators
# Accents used sparingly and only to carry meaning (pass/fail/warn).
GOOD = RGBColor(0x1E, 0x7A, 0x4D)
BAD = RGBColor(0xA3, 0x1D, 0x1D)
WARN = RGBColor(0xB0, 0x6A, 0x00)

LATIN = "Calibri"
MONO = "Consolas"

W = Inches(13.333)
H = Inches(7.5)
MARGIN = Inches(0.75)
CONTENT_W = W - 2 * MARGIN
TITLE_Y = Inches(0.55)
BODY_Y = Inches(1.62)
FOOT_Y = Inches(6.95)


def sf(run, size=14, bold=False, color=BLACK, name=LATIN, ea=None, mono=False,
       italic=False):
    """Style a run and ALWAYS set the eastAsia font.

    Playbook Section 6: Word and PowerPoint store East Asian glyphs under a
    separate font attribute (w:eastAsia / a:ea). Leave it unset and some
    builds silently substitute a serif for Chinese text while Latin text
    renders correctly. Set it on every run, not just headings.
    """
    if mono:
        name = MONO
    if ea is None:
        ea = "Microsoft YaHei" if not mono else MONO
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    latin = rPr.find(qn('a:latin'))
    if latin is None:
        latin = OxmlElement('a:latin')
        rPr.append(latin)
    latin.set('typeface', name)
    for tag in ('a:ea', 'a:cs'):
        el = rPr.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            rPr.append(el)
        el.set('typeface', ea)


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def rect(slide, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(0.75)
    s.shadow.inherit = False   # no shadows, ever
    return s


def textbox(slide, x, y, w, h, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.paragraphs[0].alignment = align
    return tb, tf


def _para(tf, first=False):
    return tf.paragraphs[0] if first else tf.add_paragraph()


def write(tf, text, first=False, size=14, bold=False, color=BLACK, mono=False,
          italic=False, align=PP_ALIGN.LEFT, space_after=0, line=None):
    p = _para(tf, first)
    p.alignment = align
    if space_after:
        p.space_after = Pt(space_after)
    if line:
        p.line_spacing = line
    r = p.add_run()
    r.text = text
    sf(r, size, bold, color, mono=mono, italic=italic)
    return p


# ------------------------------------------------------------ chrome
def chrome(slide, page, source=None, dark=False):
    """Page number bottom-right, optional source note bottom-left.

    A source line on every slide carrying a claim is standard consulting
    practice: the reader can always trace a number back.
    """
    col = MUTED if not dark else RGBColor(0x9A, 0xA8, 0xB4)
    if source:
        _, tf = textbox(slide, MARGIN, FOOT_Y, Inches(10.5), Inches(0.28))
        write(tf, source, first=True, size=9, color=col, italic=True)
    _, tf = textbox(slide, W - MARGIN - Inches(1.0), FOOT_Y, Inches(1.0),
                    Inches(0.28), align=PP_ALIGN.RIGHT)
    write(tf, str(page), first=True, size=9, color=col)


def action_title(slide, title, kicker=None):
    """The action title: a full sentence carrying the takeaway.

    kicker is the small navy label above it (the section/tracker), which
    lets a reader locate the slide in the argument without reading the body.
    """
    y = TITLE_Y
    if kicker:
        _, tf = textbox(slide, MARGIN, y, CONTENT_W, Inches(0.26))
        write(tf, kicker.upper(), first=True, size=10, bold=True, color=NAVY)
        y = y + Inches(0.32)
    _, tf = textbox(slide, MARGIN, y, CONTENT_W, Inches(0.72))
    write(tf, title, first=True, size=21, bold=True, color=NAVY, line=1.05)
    rect(slide, MARGIN, Inches(1.42), CONTENT_W, Emu(9525), fill=NAVY)


# ------------------------------------------------------------ slide types
def slide_cover(prs, kicker, title, subtitle, meta):
    s = blank(prs)
    rect(s, 0, 0, W, H, fill=NAVY_DK)
    rect(s, 0, 0, Inches(0.13), H, fill=RGBColor(0x00, 0xA6, 0xC9))
    _, tf = textbox(s, Inches(1.0), Inches(2.0), Inches(11.2), Inches(0.4))
    write(tf, kicker.upper(), first=True, size=13, bold=True,
          color=RGBColor(0x7F, 0xC4, 0xE8))
    _, tf = textbox(s, Inches(1.0), Inches(2.55), Inches(11.2), Inches(1.5))
    write(tf, title, first=True, size=38, bold=True, color=WHITE, line=1.06)
    _, tf = textbox(s, Inches(1.0), Inches(4.25), Inches(10.6), Inches(0.9))
    write(tf, subtitle, first=True, size=16, color=RGBColor(0xC9, 0xD6, 0xE0),
          line=1.3)
    rect(s, Inches(1.0), Inches(5.45), Inches(1.6), Emu(19050), fill=RGBColor(0x00, 0xA6, 0xC9))
    _, tf = textbox(s, Inches(1.0), Inches(5.75), Inches(10.6), Inches(0.9))
    for i, m in enumerate(meta):
        write(tf, m, first=(i == 0), size=11,
              color=RGBColor(0x9A, 0xA8, 0xB4), space_after=3)
    return s


def slide_divider(prs, num, title, blurb, page):
    s = blank(prs)
    rect(s, 0, 0, W, H, fill=NAVY_DK)
    _, tf = textbox(s, Inches(1.0), Inches(2.7), Inches(1.6), Inches(1.4))
    write(tf, num, first=True, size=60, bold=True,
          color=RGBColor(0x00, 0xA6, 0xC9))
    rect(s, Inches(2.5), Inches(2.85), Emu(9525), Inches(1.1),
         fill=RGBColor(0x3E, 0x5A, 0x72))
    _, tf = textbox(s, Inches(2.95), Inches(2.85), Inches(9.2), Inches(0.7))
    write(tf, title, first=True, size=30, bold=True, color=WHITE)
    _, tf = textbox(s, Inches(2.95), Inches(3.62), Inches(8.8), Inches(0.7))
    write(tf, blurb, first=True, size=14, color=RGBColor(0xB6, 0xC5, 0xD2),
          line=1.3)
    chrome(s, page, dark=True)
    return s


def slide_bullets(prs, title, bullets, page, kicker=None, source=None,
                  lead=None):
    s = blank(prs)
    action_title(s, title, kicker)
    y = BODY_Y
    if lead:
        _, tf = textbox(s, MARGIN, y, CONTENT_W, Inches(0.5))
        write(tf, lead, first=True, size=13, color=CAPTION, italic=True,
              line=1.3)
        y = y + Inches(0.62)
    # Distribute bullets across the body area instead of stacking them at a
    # fixed pitch: a fixed pitch leaves a dead band at the bottom of every
    # short slide, which reads as an unfinished layout rather than as
    # deliberate whitespace.
    avail = FOOT_Y - Inches(0.3) - y
    natural = [Inches(0.62) if len(pick_text(b)) < 95 else Inches(0.88)
               for b in bullets]
    total = sum(natural, Emu(0))
    slack = avail - total
    pad = Emu(int(slack / len(bullets))) if slack > 0 else Emu(0)
    pad = min(pad, Inches(0.42))   # cap so 2-bullet slides don't sprawl

    for i, b in enumerate(bullets):
        bold_txt, rest = (b if isinstance(b, tuple) else (None, b))
        rect(s, MARGIN + Inches(0.02), y + Inches(0.12), Inches(0.07),
             Inches(0.07), fill=NAVY)
        _, tf = textbox(s, MARGIN + Inches(0.28), y, CONTENT_W - Inches(0.28),
                        Inches(0.9))
        p = tf.paragraphs[0]
        p.line_spacing = 1.25
        if bold_txt:
            r = p.add_run(); r.text = bold_txt + "  "
            sf(r, 14, True, NAVY)
        r = p.add_run(); r.text = rest
        sf(r, 14, False, BLACK)
        y = y + natural[i] + pad
    chrome(s, page, source)
    return s


def pick_text(b):
    return b[1] if isinstance(b, tuple) else b


def slide_table(prs, title, headers, rows, page, kicker=None, source=None,
                widths=None, emphasis_col=None, lead=None):
    """Zebra table, navy header, no vertical rules (data-ink discipline)."""
    s = blank(prs)
    action_title(s, title, kicker)
    y = BODY_Y
    if lead:
        _, tf = textbox(s, MARGIN, y, CONTENT_W, Inches(0.45))
        write(tf, lead, first=True, size=13, color=CAPTION, italic=True, line=1.3)
        y = y + Inches(0.58)

    n = len(headers)
    if widths:
        total = sum(widths)
        cols = [Emu(int(CONTENT_W * w / total)) for w in widths]
    else:
        cols = [Emu(int(CONTENT_W / n))] * n

    hdr_h = Inches(0.42)
    rect(s, MARGIN, y, CONTENT_W, hdr_h, fill=NAVY)
    x = MARGIN
    for i, htxt in enumerate(headers):
        _, tf = textbox(s, x + Inches(0.12), y + Inches(0.1),
                        cols[i] - Inches(0.2), Inches(0.28))
        write(tf, htxt, first=True, size=11, bold=True, color=WHITE)
        x = x + cols[i]
    y = y + hdr_h

    # Let rows breathe into the available height rather than capping at a
    # fixed pitch and leaving a dead band beneath the table.
    avail = FOOT_Y - y - Inches(0.3)
    row_h = Emu(int(avail / max(len(rows), 1)))
    row_h = max(Inches(0.42), min(Inches(0.78), row_h))
    for ri, row in enumerate(rows):
        if ri % 2 == 1:
            rect(s, MARGIN, y, CONTENT_W, row_h, fill=FILL)
        x = MARGIN
        for ci, cell in enumerate(row):
            is_mono = cell.startswith("`") and cell.endswith("`")
            txt = cell.strip("`")
            col = BLACK
            if txt.startswith("PASS") or txt.startswith("通过"):
                col = GOOD
            elif txt.startswith("FAIL") or txt.startswith("VOID") or txt.startswith("失败") or txt.startswith("无效"):
                col = BAD
            bold = (emphasis_col is not None and ci == emphasis_col)
            _, tf = textbox(s, x + Inches(0.12), y + Inches(0.13),
                            cols[ci] - Inches(0.2), row_h - Inches(0.16))
            write(tf, txt, first=True, size=10.5, bold=bold, color=col,
                  mono=is_mono, line=1.15)
            x = x + cols[ci]
        rect(s, MARGIN, y + row_h, CONTENT_W, Emu(9525), fill=RULE)
        y = y + row_h
    chrome(s, page, source)
    return s


def slide_flow(prs, title, steps, page, kicker=None, source=None, note=None):
    """Horizontal process chevrons — the standard consulting process visual."""
    s = blank(prs)
    action_title(s, title, kicker)
    n = len(steps)
    gap = Inches(0.16)
    bw = Emu(int((CONTENT_W - gap * (n - 1)) / n))
    y = Inches(2.15)
    bh = Inches(1.65)
    x = MARGIN
    for i, (label, body) in enumerate(steps):
        shape = MSO_SHAPE.PENTAGON if i < n - 1 else MSO_SHAPE.RECTANGLE
        rect(s, x, y, bw, bh, fill=NAVY if i % 2 == 0 else RGBColor(0x2E, 0x62, 0x8F),
             shape=shape)
        _, tf = textbox(s, x + Inches(0.22), y + Inches(0.26),
                        bw - Inches(0.75), Inches(0.35))
        write(tf, label, first=True, size=13, bold=True, color=WHITE)
        _, tf = textbox(s, x + Inches(0.22), y + Inches(0.68),
                        bw - Inches(0.75), Inches(0.9))
        write(tf, body, first=True, size=10, color=RGBColor(0xD6, 0xE2, 0xEC),
              line=1.2)
        x = x + bw + gap
    if note:
        _, tf = textbox(s, MARGIN, Inches(4.35), CONTENT_W, Inches(1.6))
        write(tf, note, first=True, size=13, color=BLACK, line=1.35)
    chrome(s, page, source)
    return s


def slide_two_col(prs, title, left, right, page, kicker=None, source=None):
    """Two-panel comparison. left/right are (heading, [lines], accent)."""
    s = blank(prs)
    action_title(s, title, kicker)
    cw = Emu(int((CONTENT_W - Inches(0.35)) / 2))
    for idx, (head, lines, accent) in enumerate((left, right)):
        x = MARGIN + (cw + Inches(0.35)) * idx
        rect(s, x, BODY_Y, cw, Inches(4.9), fill=FILL)
        rect(s, x, BODY_Y, cw, Inches(0.5), fill=accent)
        _, tf = textbox(s, x + Inches(0.22), BODY_Y + Inches(0.13),
                        cw - Inches(0.44), Inches(0.3))
        write(tf, head, first=True, size=13, bold=True, color=WHITE)
        _, tf = textbox(s, x + Inches(0.22), BODY_Y + Inches(0.72),
                        cw - Inches(0.44), Inches(4.0))
        for i, ln in enumerate(lines):
            write(tf, ln, first=(i == 0), size=11.5, color=BLACK,
                  space_after=9, line=1.25)
    chrome(s, page, source)
    return s


def slide_code(prs, title, lines, page, kicker=None, source=None, note=None):
    s = blank(prs)
    action_title(s, title, kicker)
    # 11pt mono + 4pt space-after at 1.15 line spacing ~= 0.245in per line.
    # Sizing the panel from the real line height keeps the box hugging the
    # code instead of trailing an empty grey band.
    h = Inches(0.245) * len(lines) + Inches(0.5)
    rect(s, MARGIN, BODY_Y, CONTENT_W, h, fill=FILL)
    rect(s, MARGIN, BODY_Y, Inches(0.05), h, fill=NAVY)
    _, tf = textbox(s, MARGIN + Inches(0.3), BODY_Y + Inches(0.25),
                    CONTENT_W - Inches(0.6), h - Inches(0.5))
    for i, ln in enumerate(lines):
        col = BLACK
        if ln.strip().startswith("#"):
            col = CAPTION
        elif "PASS" in ln:
            col = GOOD
        elif "FAIL" in ln or "VOID" in ln:
            col = BAD
        write(tf, ln, first=(i == 0), size=11, color=col, mono=True,
              space_after=4, line=1.15)
    if note:
        _, tf = textbox(s, MARGIN, BODY_Y + h + Inches(0.3), CONTENT_W,
                        Inches(1.2))
        write(tf, note, first=True, size=13, color=BLACK, line=1.35)
    chrome(s, page, source)
    return s


def slide_big_number(prs, title, stats, page, kicker=None, source=None,
                     note=None):
    """Three-to-four KPI tiles. Number first, label under it."""
    s = blank(prs)
    action_title(s, title, kicker)
    n = len(stats)
    gap = Inches(0.3)
    bw = Emu(int((CONTENT_W - gap * (n - 1)) / n))
    y = Inches(2.1)
    x = MARGIN
    for num, label, sub in stats:
        rect(s, x, y, bw, Inches(2.15), fill=FILL)
        rect(s, x, y, bw, Inches(0.06), fill=NAVY)
        _, tf = textbox(s, x + Inches(0.25), y + Inches(0.38),
                        bw - Inches(0.5), Inches(0.8), align=PP_ALIGN.LEFT)
        write(tf, num, first=True, size=40, bold=True, color=NAVY)
        _, tf = textbox(s, x + Inches(0.25), y + Inches(1.22),
                        bw - Inches(0.5), Inches(0.3))
        write(tf, label, first=True, size=12, bold=True, color=BLACK)
        _, tf = textbox(s, x + Inches(0.25), y + Inches(1.55),
                        bw - Inches(0.5), Inches(0.5))
        write(tf, sub, first=True, size=10, color=CAPTION, line=1.2)
        x = x + bw + gap
    if note:
        _, tf = textbox(s, MARGIN, Inches(4.65), CONTENT_W, Inches(1.5))
        write(tf, note, first=True, size=13, color=BLACK, line=1.35)
    chrome(s, page, source)
    return s


def slide_quote(prs, quote, attrib, page, kicker=None):
    s = blank(prs)
    rect(s, 0, 0, W, H, fill=NAVY_DK)
    rect(s, MARGIN, Inches(2.4), Inches(0.06), Inches(2.4),
         fill=RGBColor(0x00, 0xA6, 0xC9))
    if kicker:
        _, tf = textbox(s, Inches(1.25), Inches(2.0), Inches(10.5), Inches(0.3))
        write(tf, kicker.upper(), first=True, size=10, bold=True,
              color=RGBColor(0x7F, 0xC4, 0xE8))
    _, tf = textbox(s, Inches(1.25), Inches(2.45), Inches(10.6), Inches(2.2))
    write(tf, quote, first=True, size=27, bold=True, color=WHITE, line=1.25)
    _, tf = textbox(s, Inches(1.25), Inches(4.95), Inches(10.6), Inches(0.4))
    write(tf, attrib, first=True, size=12, color=RGBColor(0x9A, 0xA8, 0xB4))
    chrome(s, page, dark=True)
    return s
